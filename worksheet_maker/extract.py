"""업로드 파일 준비.

- PDF·이미지는 AI가 직접 보도록 첨부 파일(Attachment)로 넘긴다. 스캔본이나 사진도 읽을 수 있다.
- DOCX·HWPX·TXT는 글자만 뽑아서 넘긴다.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from .llm import Attachment

ATTACH_TYPES = {
    "pdf": "application/pdf",
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
}
TEXT_TYPES = ["docx", "hwpx", "txt", "md"]
SUPPORTED = list(ATTACH_TYPES) + TEXT_TYPES + ["hwp"]  # hwp는 안내 메시지를 보여주려고 받는다
VISUAL_TYPES = list(ATTACH_TYPES)  # '학습지 따라 만들기'에서 받는 형식


class ExtractError(Exception):
    pass


def prepare(filename: str, data: bytes) -> str | Attachment:
    """파일을 AI에 넘길 형태로 바꾼다: 첨부(Attachment) 또는 텍스트(str)."""
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext in ATTACH_TYPES:
        return Attachment(name=filename, mime_type=ATTACH_TYPES[ext], data=data)
    if ext == "docx":
        return _docx(data)
    if ext == "hwpx":
        return _hwpx(data)
    if ext == "hwp":
        raise ExtractError("HWP 파일은 읽을 수 없습니다. 한글에서 HWPX 또는 PDF로 저장한 뒤 올려 주세요.")
    if ext in ("txt", "md"):
        for enc in ("utf-8-sig", "cp949"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        raise ExtractError("텍스트 파일의 인코딩을 알 수 없습니다.")
    raise ExtractError(f"지원하지 않는 파일 형식입니다: .{ext}")


def prepare_many(uploads: list[tuple[str, bytes]]) -> tuple[str, list[Attachment]]:
    texts, files = [], []
    for name, data in uploads:
        item = prepare(name, data)
        if isinstance(item, Attachment):
            files.append(item)
        else:
            texts.append(f"[{name}]\n{item}")
    return "\n\n".join(texts), files


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _hwpx(data: bytes) -> str:
    # HWPX는 ZIP 안의 Contents/section0.xml, section1.xml … 에 본문이 들어 있다.
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        sections = sorted(
            (n for n in z.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n)),
            key=lambda n: int(re.search(r"\d+", n).group()),
        )
        if not sections:
            raise ExtractError("HWPX 본문을 찾지 못했습니다.")
        lines = []
        for name in sections:
            root = ElementTree.fromstring(z.read(name))
            for p in root.iter():
                if _local(p.tag) != "p":
                    continue
                # 표 안의 문단이 중복되지 않도록 p > run > t 만 모은다.
                text = "".join(
                    "".join(t.itertext())
                    for run in p if _local(run.tag) == "run"
                    for t in run if _local(t.tag) == "t"
                )
                lines.append(text)
        return "\n".join(lines)
